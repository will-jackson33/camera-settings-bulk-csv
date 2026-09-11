"""Read an edit sheet - the spreadsheet a person hands over - into rows of column name to text.

An .xlsx is read with openpyxl; a .csv or .tsv with the standard library. The first row is the
header. Every cell comes back as the text a person would have typed: whole numbers without a
trailing .0, booleans as True or False, dates as ISO text, blanks as ''. What the columns mean is
camerasettings_edit's business; this module only makes the sheet readable.
"""

from __future__ import annotations

import csv
import datetime
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

XLSX_SUFFIXES = (".xlsx", ".xlsm")
TEXT_SUFFIXES = {".csv": ",", ".tsv": "\t", ".txt": "\t"}


class SheetError(ValueError):
    """The sheet cannot be read as a table: wrong kind of file, no header, or a repeated column."""


@dataclass
class SheetRow:
    number: int  # the spreadsheet row number, header being row 1, for messages
    values: dict[str, str]


@dataclass
class EditSheet:
    source: str
    columns: list[str]
    rows: list[SheetRow]


def camerasettings_sheet_cellText(cell: object) -> str:
    """The text a person would have typed into the cell. bool before int: bool is an int in Python."""
    if cell is None:
        return ""
    if isinstance(cell, bool):
        return "True" if cell else "False"
    if isinstance(cell, int):
        return str(cell)
    if isinstance(cell, float):
        return str(int(cell)) if cell.is_integer() else repr(cell)
    if isinstance(cell, datetime.datetime):
        return cell.isoformat(sep=" ")
    if isinstance(cell, datetime.date):
        return cell.isoformat()
    return str(cell).strip()


def _rows_from_xlsx(path: Path, worksheet: str | None) -> list[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if worksheet is not None and worksheet not in workbook.sheetnames:
            raise SheetError(f"{path.name} has no worksheet named {worksheet}; it has {', '.join(workbook.sheetnames)}")
        sheet = workbook[worksheet] if worksheet is not None else workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _decode(data: bytes) -> str:
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        return data.decode("utf-16")
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("cp1252")


def _rows_from_text(path: Path, delimiter: str) -> list[list[object]]:
    text = _decode(path.read_bytes())
    return [list(row) for row in csv.reader(text.splitlines(), delimiter=delimiter)]


def _table(raw: list[list[Any]], source: str) -> EditSheet:
    if not raw:
        raise SheetError(f"{source} is empty: no header row")
    names = [camerasettings_sheet_cellText(cell) for cell in raw[0]]
    columns = [name for name in names if name]
    if not columns:
        raise SheetError(f"{source} has no column names in its first row")
    repeated = sorted({name for name in columns if columns.count(name) > 1})
    if repeated:
        raise SheetError(f"{source} names a column twice: {', '.join(repeated)}")
    rows: list[SheetRow] = []
    for number, cells in enumerate(raw[1:], 2):
        texts = [camerasettings_sheet_cellText(cell) for cell in cells]
        values = {name: (texts[i] if i < len(texts) else "") for i, name in enumerate(names) if name}
        if any(values.values()):
            rows.append(SheetRow(number, values))
    return EditSheet(source, columns, rows)


def camerasettings_sheet_read(path: Path, worksheet: str | None = None) -> EditSheet:
    """Read an .xlsx, .xlsm, .csv, .tsv or .txt edit sheet.

    Blank-named columns are ignored (Excel leaves them at the right edge), blank rows are skipped,
    and a row's missing trailing cells read as blank. Raises SheetError for a file kind it cannot
    read, a missing file, an empty sheet, or a column name that appears twice.
    """
    if not path.exists():
        raise SheetError(f"{path} does not exist")
    suffix = path.suffix.lower()
    if suffix in XLSX_SUFFIXES:
        return _table(_rows_from_xlsx(path, worksheet), path.name)
    if suffix in TEXT_SUFFIXES:
        return _table(_rows_from_text(path, TEXT_SUFFIXES[suffix]), path.name)
    raise SheetError(f"{path.name}: an edit sheet is .xlsx, .csv or .tsv, not {suffix or 'a file without an extension'}")
